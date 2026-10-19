import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as ReportLabImage, PageBreak

# Resim Dosyalarının Adları (Proje klasöründeki dosyalarla eşleşmeli)
IMG_DEADLOCK = "resim1.jpeg"
IMG_DYNAMIC = "resim2.jpeg"
IMG_RESOLVED = "resim3.jpeg"

def generate_word():
    doc = Document()
    
    # Kapak Başlığı
    title = doc.add_heading("Deadlock Simulator & Banker's Algorithm", level=0)
    p_meta = doc.add_paragraph("Kapsamlı Jüri Sunum Raporu ve Resimli Kod Analizi (7 Dakikalık Sunum Rehberi)\nHazırlayan: Ege Akdeniz | İşletim Sistemleri")
    p_meta.runs[0].font.color.rgb = RGBColor(90, 90, 90)
    doc.add_paragraph("=" * 80)

    # 1. BÖLÜM
    doc.add_heading("1. Proje Amacı ve Coffman Kuralları Analizi", level=1)
    doc.add_paragraph(
        "Modern işletim sistemlerinde birden fazla sürecin kısıtlı kaynakları aynı anda talep etmesi döngüsel kilitlenmelere (Deadlock) yol açar. "
        "Bir deadlock durumunun oluşması için 4 temel Coffman koşulunun aynı anda gerçekleşmesi gerekir:\n"
        "• Karşılıklı Dışlama (Mutual Exclusion): Kaynak aynı anda yalnızca tek süreç tarafından kullanılabilir.\n"
        "• Tut ve Bekle (Hold and Wait): Süreç bir kaynağı tutarken diğerini bekler.\n"
        "• Kesilmeme (No Preemption): Kaynak, süreç tamamlanmadan zorla geri alınamaz.\n"
        "• Döngüsel Bekleme (Circular Wait): Süreçler arasında dairesel bekleme zinciri oluşur.\n"
        "Geliştirilen bu yazılım; Banker's Algoritması ile güvenlik denetimi yapar, RAG grafiği çizer ve bu 4 kuralı kırarak sistemi kurtarır."
    )

    # 2. BÖLÜM
    doc.add_heading("2. Modüler Yazılım Mimarisi (Kod Ayrımı)", level=1)
    doc.add_paragraph(
        "Yazılım üç bağımsız katmandan meydana gelir:\n"
        "1. models.py (Engine Katmanı): Durum matrisleri (Allocation, Max Need, Available, Priority, Cost), "
        "Banker's Algoritması, adım adım yürütücü durum üreteci ve kurban seçim algoritmalarını barındırır.\n"
        "2. gui.py (Arayüz & Grafik Katmanı): Tkinter ttk bileşenleri, Matplotlib FigureCanvasTkAgg ve NetworkX tabanlı yönlü grafik çizimlerini yönetir.\n"
        "3. main.py (Başlatıcı): Uygulama döngüsünü başlatan giriş noktasıdır."
    )

    # 3. BÖLÜM - RESİM 1 VE DEADLOCK ANALİZİ
    doc.add_heading("3. Ekran Analizi 1: Kilitlenme (Deadlock) Durumu ve RAG Grafiği", level=1)
    doc.add_paragraph(
        "Aşağıdaki ekran görüntüsünde 3 Süreç (P0, P1, P2) ve 2 Kaynak (R0, R1) içeren bir Deadlock senaryosu görülmektedir. "
        "Available vektörü [0, 0] olup sistemde boştaki kaynak kalmamıştır."
    )
    
    if os.path.exists(IMG_DEADLOCK):
        doc.add_picture(IMG_DEADLOCK, width=Inches(6.0))
        caption = doc.add_paragraph("Şekil 1: Kilitlenme (Deadlock) Anı - Kaynak Tahsis Grafiğinde Döngüsel Bekleme")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].font.italic = True
        caption.runs[0].font.size = Pt(9)
    else:
        doc.add_paragraph(f"[UYARI: {IMG_DEADLOCK} bulunamadı. Lütfen görseli proje klasörüne ekleyin.]")

    doc.add_paragraph(
        "Koddaki Çizim Mantığı (draw_graph Fonksiyonu):\n"
        "• Mavi Çizgiler (Tahsis - Allocation): R -> P yönündeki oklardır. Kodda 'allocation[i][j] > 0' olduğunda çizilir. "
        "R0->P0, R1->P1 ve P2'nin hem R0 hem R1'i tuttuğu görülmektedir.\n"
        "• Turuncu Çizgiler (Talep/Bekleme - Request): P -> R yönündeki oklardır. Kodda 'need > 0 and available[j] < need' "
        "şartı sağlandığında çizilir. P0->R1 ve P1->R0 okları çapraz bir döngüsel bekleme (Circular Wait) oluşturmuştur.\n"
        "• Jüri Notu: Available [0, 0] olduğu için hiçbir süreç ihtiyacını tamamlayamaz ve sistem tamamen kilitlenir."
    )

    # 4. BÖLÜM - RESİM 2 VE BOYUTLANDIRMA
    doc.add_heading("4. Ekran Analizi 2: Dinamik Boyutlandırma ve Adım Motoru", level=1)
    doc.add_paragraph(
        "Sistem esnek bir yapıya sahiptir; kullanıcı arayüz üzerinden süreç ve kaynak sayılarını dinamik olarak değiştirebilir."
    )

    if os.path.exists(IMG_DYNAMIC):
        doc.add_picture(IMG_DYNAMIC, width=Inches(6.0))
        caption = doc.add_paragraph("Şekil 2: Dinamik Boyutlandırma ve Sistem Durum Logu")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].font.italic = True
        caption.runs[0].font.size = Pt(9)
    else:
        doc.add_paragraph(f"[UYARI: {IMG_DYNAMIC} bulunamadı.]")

    doc.add_paragraph(
        "Koddaki Dinamik Boyutlandırma ve Adım Yürütme Detayları:\n"
        "• set_dimensions(new_p, new_r): Kullanıcı 'Uygula' butonuna bastığında bellek matrisleri yeniden tahsis edilir.\n"
        "• create_matrix_inputs(): Eski kutucuklar 'widget.destroy()' ile yok edilip yeni boyuta göre dinamik grid açılır.\n"
        "• Adım Motoru Renk Kodları: İnceleme sırasındaki süreç sarı (#f1c40f), kaynak alan ve serbest bırakan yeşil (#2ecc71), "
        "kaynak yetersizliğinden bekletilen süreç kırmızı (#e74c3c) renkle vurgulanır."
    )

    # 5. BÖLÜM - RESİM 3 VE KURBAN SEÇİMİYLE ÇÖZÜM
    doc.add_heading("5. Ekran Analizi 3: Akıllı Kurban Seçimi ve Kilitlenmenin Çözümü", level=1)
    doc.add_paragraph(
        "Deadlock oluştuktan sonra 'Akıllı Kilitlenme Çözümü' panelinden strateji uygulanmış ve sistem kurtarılmıştır."
    )

    if os.path.exists(IMG_RESOLVED):
        doc.add_picture(IMG_RESOLVED, width=Inches(6.0))
        caption = doc.add_paragraph("Şekil 3: P2 Sürecinin Feda Edilmesi Sonucu Kilitlenmenin Çözülmesi")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].font.italic = True
        caption.runs[0].font.size = Pt(9)
    else:
        doc.add_paragraph(f"[UYARI: {IMG_RESOLVED} bulunamadı.]")

    doc.add_paragraph(
        "Çözüm Mekanizmasının Kod Analizi:\n"
        "• resolve_deadlock_strategy() fonksiyonunda 'Victim: Max Allocation' veya 'Resource Preemption' stratejisi seçilmiştir.\n"
        "• Kod 'total_allocated = [sum(allocation[i])]' listesi üzerinden en çok kaynak bloke eden P2 sürecini kurban seçmiştir.\n"
        "• P2'nin elindeki kaynaklar 'allocation[2] = [0, 0]' yapılarak havuza aktarılmış ve Available vektörü [1, 1] olmuştur.\n"
        "• Sonuç: Şekil 3'te görüldüğü gibi P2'nin tüm bağları kopmuş, açığa çıkan [1, 1] kaynağı sayesinde P0 ve P1'in talepleri karşılanmış ve sistem güvenli duruma geçmiştir."
    )

    # 6. BÖLÜM - SUNUM REHBERİ
    doc.add_heading("6. 7 Dakikalık Jüri Sunum Konuşma Rehberi", level=1)
    doc.add_paragraph(
        "• 0:00 - 1:00 (Giriş): Deadlock tanımı, projenin modüler amacı ve 4 Coffman şartı aktarılır.\n"
        "• 1:00 - 2:00 (Mimari): models.py (mantık), gui.py (arayüz) ve main.py yapısı anlatılır.\n"
        "• 2:00 - 3:30 (Banker Algoritması): Need = Max - Alloc ve Work vektörü döngüsü açıklanır. Zaman karmaşıklığının O(m * n^2) olduğu belirtilir.\n"
        "• 3:30 - 4:45 (Resim 1 & Resim 2): Şekil 1'deki döngüsel oklar, mavi/turuncu çizgi mantığı ve Şekil 2'deki dinamik grid yapısı gösterilir.\n"
        "• 4:45 - 6:00 (Resim 3 & Çözüm): Şekil 3 üzerinden P2'nin nasıl kurban seçildiği ve Available vektörünün [1, 1] yapılarak kilitlenmenin nasıl çözüldüğü aktarılır.\n"
        "• 6:00 - 7:00 (Stratejiler & Kapanış): 4 kurban seçme algoritması (Max Alloc, Priority, Cost, Preemption) özetlenerek sunum tamamlanır."
    )

    doc.save("Deadlock_Proje_Raporu.docx")
    print("✓ Resimli Word Raporu (Deadlock_Proje_Raporu.docx) başarıyla oluşturuldu.")


def generate_pdf():
    pdf_filename = "Deadlock_Proje_Raporu.pdf"
    doc = SimpleDocTemplate(pdf_filename, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=15, leading=18, textColor=colors.HexColor('#1a237e'))
    h2_style = ParagraphStyle('DocH2', parent=styles['Heading2'], fontSize=11, leading=14, textColor=colors.HexColor('#0d47a1'), spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle('DocBody', parent=styles['Normal'], fontSize=8.5, leading=11.5, spaceAfter=4)
    caption_style = ParagraphStyle('DocCap', parent=styles['Italic'], fontSize=7.5, leading=9, alignment=1, textColor=colors.HexColor('#555555'), spaceAfter=4)

    story = []
    story.append(Paragraph("Deadlock Simulator & Banker's Algorithm (Pro Edition)", title_style))
    story.append(Paragraph("Kapsamli Teknik Rapor, Resimli Kod Analizi ve 7 Dakikalik Sunum Rehberi | Ege Akdeniz", body_style))
    story.append(Spacer(1, 6))

    # Bölüm 1 & 2
    story.append(Paragraph("1. Proje Ozeti ve Moduler Yapi", h2_style))
    story.append(Paragraph("Bu proje; coklu gorev isletim sistemlerindeki deadlock (kilitlenme) durumunu analiz eden, Banker's Algoritmasi ile guvenlik denetimi yapan ve Resource Allocation Graph (RAG) uzerinde gorsellestiren masaustu yazilimidir. Kod uc ana mod 카mandan olusur: models.py (hesaplama ve cozum motoru), gui.py (arayuz ve NetworkX grafigi), main.py (baslatici).", body_style))

    # Resim 1 ve Açıklaması
    story.append(Paragraph("2. Resim 1 Analizi: Deadlock Durumu ve RAG Grafigi", h2_style))
    story.append(Paragraph("Mavi cizgiler (Allocation > 0) tahsis edilen kaynaklari (R -> P), turuncu cizgiler (Need > 0 ve Available < Need) talep edilen kaynaklari (P -> R) temsil eder. Available [0, 0] iken P0 ve P1 capraz bekleme olusturmus, P2 ise kaynaklari bloke etmistir.", body_style))
    if os.path.exists(IMG_DEADLOCK):
        story.append(ReportLabImage(IMG_DEADLOCK, width=5.5*72, height=3.0*72))
        story.append(Paragraph("Sekil 1: Kilitlenme Durumu (RAG Grafiginde Kapali Dongu)", caption_style))

    # Resim 2 ve Açıklaması
    story.append(Paragraph("3. Resim 2 Analizi: Dinamik Boyutlandirma ve Adim Motoru", h2_style))
    story.append(Paragraph("Surec ve kaynak sayilari dinamik degistirilebilir. Adim motorunda kontrol edilen proses sari, kaynak alan yesil, engellenen proses kirmizi renkle vurgulanir.", body_style))
    if os.path.exists(IMG_DYNAMIC):
        story.append(ReportLabImage(IMG_DYNAMIC, width=5.5*72, height=3.0*72))
        story.append(Paragraph("Sekil 2: Dinamik Boyutlandirma ve Sistem Durum Logu", caption_style))

    # Resim 3 ve Açıklaması
    story.append(Paragraph("4. Resim 3 Analizi: Kilitlenmenin Cozulmesi", h2_style))
    story.append(Paragraph("Akilli cozum stratejisi ile en cok kaynak tutan P2 kurban edilmis, kaynaklari serbest birakilarak Available [1, 1] yapilmis ve sistem guvenli duruma gecmistir.", body_style))
    if os.path.exists(IMG_RESOLVED):
        story.append(ReportLabImage(IMG_RESOLVED, width=5.5*72, height=3.0*72))
        story.append(Paragraph("Sekil 3: P2 Surecinin Feda Edilmesiyle Kilitlenmenin Cozulmesi", caption_style))

    # Bölüm 5
    story.append(Paragraph("5. 7 Dakikalik Sunum Akisi", h2_style))
    story.append(Paragraph("0:00-1:00 Giris ve Coffman kurallari | 1:00-2:00 Moduler mimari | 2:00-3:30 Banker Algoritmasi | 3:30-4:45 Sekil 1 ve Sekil 2 grafikleri | 4:45-6:00 Sekil 3 kilitlenmenin cozulusu | 6:00-7:00 Cozum stratejileri ve Soru-Cevap.", body_style))

    doc.build(story)
    print("✓ Resimli PDF Raporu (Deadlock_Proje_Raporu.pdf) başarıyla oluşturuldu.")


if __name__ == "__main__":
    generate_word()
    generate_pdf()