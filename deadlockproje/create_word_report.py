import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Klasördeki resimlerin isimleri (WhatsApp görsellerinin adları veya resim1.jpeg vb.)
IMG_1_CANDIDATES = ["WhatsApp Image 2026-10-19 at 12.49.00.jpeg", "resim1.jpeg"]
IMG_2_CANDIDATES = ["WhatsApp Image 2026-10-19 at 12.49.00 (1).jpeg", "resim2.jpeg"]
IMG_3_CANDIDATES = ["WhatsApp Image 2026-10-19 at 12.49.00 (2).jpeg", "resim3.jpeg"]

def find_image(candidates):
    for name in candidates:
        if os.path.exists(name):
            return name
    return None

def add_custom_box(doc, title_text, body_text, r, g, b):
    p = doc.add_paragraph()
    run_title = p.add_run(f"{title_text}\n")
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(r, g, b)
    run_body = p.add_run(body_text)
    run_body.font.size = Pt(10)

def generate_comprehensive_word():
    doc = Document()

    # Ana Başlık
    title = doc.add_heading("DEADLOCK SIMULATOR & BANKER'S ALGORITHM", level=0)
    sub = doc.add_paragraph(
        "20 Dakikalık Jüri Sunum Rehberi, Satır Satır Kod Ezber Kılavuzu ve Ekran Görüntüsü Analizi\n"
        "Hazırlayan: Ege Akdeniz | İşletim Sistemleri Dönem Projesi"
    )
    sub.runs[0].font.color.rgb = RGBColor(80, 80, 80)
    doc.add_paragraph("=" * 85)

    # 1. ZAMAN PLANI
    doc.add_heading("1. 20 Dakikalık Sunumun Zaman ve Strateji Planı", level=1)
    timeline = [
        ("00:00 – 03:00 | Giriş ve Coffman Koşulları", "Deadlock tanımı, 4 temel şart (Mutual Exclusion, Hold & Wait, No Preemption, Circular Wait)."),
        ("03:00 – 06:00 | Mimari ve Kütüphaneler", "main.py, models.py, gui.py ayrımı, FigureCanvasTkAgg ve nx.DiGraph kullanımı."),
        ("06:00 – 09:30 | Veri Yapıları ve Matrisler", "models.py içindeki __init__, set_dimensions, List Comprehension ve Shallow Copy önleme."),
        ("09:30 – 13:00 | Banker's Algorithm ve Adım Motoru", "run_bankers(), Need = Max - Allocation, all() fonksiyonu ve O(m * n^2) zaman karmaşıklığı."),
        ("13:00 – 16:00 | RAG Grafiği ve Resim 1-2 Analizi", "draw_graph() koordinat matematiği, Mavi/Turuncu ok şartları, dinamik boyutlandırma."),
        ("16:00 – 18:30 | Akıllı Çözüm Stratejileri ve Resim 3", "resolve_deadlock_strategy(), 4 kurban seçimi algoritması, P2'nin kaynaklarının geri alınması."),
        ("18:30 – 20:00 | Raporlama ve Soru-Cevap", "export_report(), filedialog, UTF-8 dosya yazımı ve jüri sorularının yanıtlanması.")
    ]
    for time_slot, desc in timeline:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{time_slot}: ").bold = True
        p.add_run(desc)

    # 2. TEORİK TEMELLER
    doc.add_heading("2. Teorik Temeller ve 4 Coffman Şartı (00:00 – 03:00)", level=1)
    doc.add_paragraph(
        "Çoklu görev işletim sistemlerinde süreçler (P) kaynaklara (R) ihtiyaç duyar. Bir Deadlock oluşması için "
        "Edward G. Coffman tarafından tanımlanan şu 4 koşulun aynı anda sağlanması şarttır:\n"
        "1. Mutual Exclusion (Karşılıklı Dışlama): Kaynak aynı anda tek bir süreç tarafından kullanılabilir.\n"
        "2. Hold and Wait (Tut ve Bekle): Süreç elinde bir kaynak tutarken başka bir kaynağı bekler.\n"
        "3. No Preemption (Kesilmeme): Kaynak, süreç işini bitirmeden zorla elinden alınamaz.\n"
        "4. Circular Wait (Döngüsel Bekleme): P0 -> R1 -> P1 -> R0 -> P0 şeklinde kapalı bir bekleme zinciri oluşur."
    )
    add_custom_box(
        doc,
        "🎙️ HOCAYA SÖYLENECEK AÇILIŞ KONUŞMASI:",
        "\"Hocam merhabalar. Projemiz işletim sistemlerindeki ölümcül kilitlenmeleri hem matris düzeyinde hem de yönlü çizge "
        "(RAG) üzerinde simüle eden modüler bir yazılımdır. Sistemimiz önce Dijkstra'nın Banker Algoritması ile güvenli durumu "
        "denetlemekte, kilitlenme varsa 4 akıllı kurtarma stratejisiyle 'No Preemption' ve 'Circular Wait' şartlarını kırarak sistemi kurtarmaktadır.\"",
        30, 64, 175
    )

    # 3. KÜTÜPHANELER VE MİMARİ
    doc.add_heading("3. Kütüphaneler ve Modüler Mimari (03:00 – 06:00)", level=1)
    libs = [
        ("import tkinter as tk", "main.py ve gui.py içinde ana pencere (tk.Tk), metin log ekranı (tk.Text) ve yerleşim için kullanıldı."),
        ("from tkinter import ttk", "Modern Combobox, Spinbox ve Dark/Light tema yönetimi (ttk.Style) için kullanıldı."),
        ("messagebox, filedialog", "Uyarı pencereleri açmak ve .txt raporunu kaydederken işletim sistemi penceresi açmak için kullanıldı."),
        ("import networkx as nx", "nx.DiGraph() ile yönlü çizge oluşturup süreç ve kaynak düğümlerini modellemek için kullanıldı."),
        ("matplotlib.pyplot as plt", "plt.subplots() ile grafiğin çizileceği figür (self.fig) ve eksen (self.ax) tuvalini oluşturur."),
        ("FigureCanvasTkAgg", "Matplotlib grafiğini ayrı pencere açmadan doğrudan Tkinter içine gömmek için kullanıldı."),
        ("import datetime", "models.py içinde üretilen sistem raporuna tarih ve saat damgası vurmak için kullanıldı.")
    ]
    for lib_name, lib_desc in libs:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(f"{lib_name}: ").bold = True
        p.add_run(lib_desc)

    add_custom_box(
        doc,
        "❓ HOCA SORUSU 1: Neden plt.show() yerine FigureCanvasTkAgg kullandınız?",
        "Cevap: Normal plt.show() bağımsız pencere açar ve Tkinter döngüsünü dondurur (blocking). Biz arayüz donmadan aynı pencere "
        "içinde canlı animasyon yapmak için FigureCanvasTkAgg köprüsünü kullandık.",
        180, 83, 9
    )

    # 4. MODELS.PY VE BELLEK YÖNETİMİ
    doc.add_heading("4. models.py Veri Yapıları ve Bellek Yönetimi (06:00 – 09:30)", level=1)
    doc.add_paragraph(
        "models.py içindeki DeadlockEngine sınıfı tamamen bağımsız algoritma çekirdeğidir:\n"
        "• self.allocation: P x R boyutunda tahsis matrisi.\n"
        "• self.max_need: P x R boyutunda maksimum ihtiyaç matrisi.\n"
        "• self.available: 1 x R boyutunda boştaki kaynak vektörü.\n"
        "• self.priorities ve self.costs: Süreçlerin öncelik ve sonlandırma maliyet vektörleri."
    )
    add_custom_box(
        doc,
        "❓ HOCA SORUSU 2: set_dimensions içinde neden [[0]*new_r]*new_p yazmadınız da List Comprehension kullandınız?",
        "Cevap: Python'da çarpma operatörüyle 2 boyutlu liste oluşturulursa Shallow Copy (Sığ Kopyalama) oluşur; tüm satırlar RAM'de aynı "
        "adresi gösterir ve P0 değişince P1 de değişir. Her satıra bağımsız bellek tahsis etmek için [[0 for _ in range(new_r)] for _ in range(new_p)] kullandık.",
        180, 83, 9
    )

    # 5. BANKER ALGORİTMASI
    doc.add_heading("5. Banker's Algorithm (run_bankers) Satır Satır Analizi (09:30 – 13:00)", level=1)
    doc.add_paragraph(
        "1. work = list(self.available): Orijinal available listesini bozmamak için kopyası alınır.\n"
        "2. need[i][j] = max_need[i][j] - allocation[i][j]: Her sürecin kalan kaynak ihtiyacı hesaplanır.\n"
        "3. all(need[i][j] <= work[j] for j in range(num_resources)): Sürecin tüm kaynak ihtiyaçları eldeki work vektöründen küçük veya eşit mi diye bakılır.\n"
        "4. Koşul sağlanırsa work[j] += allocation[i][j] ile süreç kaynakları kullanıp sisteme geri iade eder ve finish[i] = True olur.\n"
        "5. Zaman Karmaşıklığı (Time Complexity): Süreç sayısı n, kaynak sayısı m iken en kötü durumda O(m * n^2) karmaşıklığındadır."
    )

    # 6. RESİM 1 VE RAG GRAFİĞİ
    doc.add_heading("6. RAG Grafiği Kod Analizi ve Ekran Görüntüsü 1 (13:00 – 14:30)", level=1)
    img1 = find_image(IMG_1_CANDIDATES)
    if img1:
        doc.add_picture(img1, width=Inches(6.0))
        c1 = doc.add_paragraph("Şekil 1: 3 Süreç ve 2 Kaynak ile Oluşan Döngüsel Kilitlenme (Deadlock)")
        c1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "draw_graph() Kenar Kuralları:\n"
        "• Mavi Çizgiler (Tahsis - R -> P): Kodda 'if allocation[i][j] > 0' şartıyla çizilir.\n"
        "• Turuncu Çizgiler (Talep - P -> R): Kodda 'if need > 0 and available[j] < need' şartıyla çizilir.\n"
        "Şekil 1'de Available = [0, 0] olduğu için P0 ve P1'in kalan ihtiyaçları karşılanamaz; turuncu talep okları kapalı bir döngü oluşturarak Deadlock meydana getirir."
    )

    # 7. RESİM 2 VE DİNAMİK BOYUTLANDIRMA
    doc.add_heading("7. Dinamik Boyutlandırma ve Ekran Görüntüsü 2 (14:30 – 16:00)", level=1)
    img2 = find_image(IMG_2_CANDIDATES)
    if img2:
        doc.add_picture(img2, width=Inches(6.0))
        c2 = doc.add_paragraph("Şekil 2: Dinamik Süreç/Kaynak Boyutlandırma ve Sistem Durum Logu")
        c2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "create_matrix_inputs() içinde 'for widget in input_group.winfo_children(): widget.destroy()' komutu kullanılarak "
        "eski giriş kutuları bellekten silinir ve yeni süreç/kaynak boyutuna göre dinamik olarak yeniden dizilir."
    )

    # 8. RESİM 3 VE AKILLI ÇÖZÜM
    doc.add_heading("8. Akıllı Çözüm Stratejileri ve Ekran Görüntüsü 3 (16:00 – 18:30)", level=1)
    img3 = find_image(IMG_3_CANDIDATES)
    if img3:
        doc.add_picture(img3, width=Inches(6.0))
        c3 = doc.add_paragraph("Şekil 3: P2 Sürecinin Kurban Seçilip Kaynaklarının Alınması Sonrası Kilitlenmenin Çözülmesi")
        c3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "resolve_deadlock_strategy() içindeki 4 algoritma:\n"
        "1. Victim: Max Allocation -> En çok kaynak bloke eden süreci (sum(allocation[i])) seçer.\n"
        "2. Lowest Priority -> Öncelik sayısı en büyük (en önemsiz) olan süreci (max(priorities)) seçer.\n"
        "3. Lowest Termination Cost -> Sonlandırma maliyeti en düşük olanı (min(costs)) seçer.\n"
        "4. Resource Preemption -> Seçilen sürecin kaynaklarını zorla alıp available vektörüne ekler.\n"
        "Şekil 3'te en çok kaynak tutan P2 ([1, 1]) kurban seçilmiş, kaynakları geri alınarak Available = [1, 1] yapılmış ve P0 ile P1'in turuncu bekleme okları çözülerek sistem güvenli duruma geçmiştir."
    )

    output_name = "Deadlock_20Dk_Juri_Sunum_Rehberi.docx"
    doc.save(output_name)
    print(f"✓ {output_name} başarıyla oluşturuldu!")

if __name__ == "__main__":
    generate_comprehensive_word()